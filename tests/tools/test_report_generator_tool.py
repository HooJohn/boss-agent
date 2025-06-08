import os
import pytest
from unittest.mock import MagicMock, patch
from boss_agent.tools.report_generator_tool import ReportGeneratorTool
from boss_agent.utils import WorkspaceManager
from boss_agent.llm.base import LLMClient
from boss_agent.tools.base import ToolImplOutput
from docx import Document
from openpyxl import Workbook
from PyPDF2 import PdfWriter, PageObject

@pytest.fixture
def workspace_manager(tmp_path):
    """Fixture for a temporary workspace manager with diverse file types."""
    manager = WorkspaceManager(tmp_path)
    finance_dir = tmp_path / "finance"
    finance_dir.mkdir()
    sales_dir = tmp_path / "sales"
    sales_dir.mkdir()

    # CSV file
    (finance_dir / "income-statement-2024-Q2.csv").write_text("Date,Revenue,Expenses\n2024-04-01,1000,500")
    
    # TXT file
    (finance_dir / "balance-sheet-2023.txt").write_text("Assets: 5000, Liabilities: 2000")

    # JSON file
    (sales_dir / "sales-summary-2024-07.json").write_text('{"total_sales": 50000, "units_sold": 120}')

    # DOCX file
    doc = Document()
    doc.add_paragraph("This is a test document for sales report 2024.")
    doc.save(sales_dir / "sales-report-2024.docx")

    # XLSX file
    wb = Workbook()
    ws = wb.active
    ws.title = "Q3 Sales"
    ws.append(["Product", "Units Sold"])
    ws.append(["Product A", 150])
    wb.save(sales_dir / "sales-details-2024-Q3.xlsx")

    # PDF file (simple one-page)
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    # This is a simplified way to create a PDF for testing purposes.
    # In a real scenario, you might use a library like reportlab to add text.
    with open(finance_dir / "financial-overview-2023.pdf", "wb") as f:
        writer.write(f)

    return manager

@pytest.fixture
def llm_client():
    """Fixture for a mock LLM client."""
    client = MagicMock()
    client.chat_completion.return_value.completion = "This is the generated report."
    return client

@pytest.fixture
def report_generator_tool(workspace_manager, llm_client):
    """Fixture for the ReportGeneratorTool."""
    return ReportGeneratorTool(workspace_manager=workspace_manager, client=llm_client)

class TestReportGeneratorTool:
    def test_find_relevant_files_success(self, report_generator_tool):
        """Test that relevant files are found correctly."""
        files = report_generator_tool._find_relevant_files(
            department="finance",
            time_dimension="2024-Q2",
            report_type="income_statement"
        )
        assert len(files) == 1
        assert "income-statement-2024-Q2.csv" in files[0]

    def test_find_relevant_files_no_match(self, report_generator_tool):
        """Test that no files are returned when there is no match."""
        files = report_generator_tool._find_relevant_files(
            department="finance",
            time_dimension="2025",
            report_type="non_existent_report"
        )
        assert len(files) == 0

    def test_find_relevant_files_department_not_found(self, report_generator_tool):
        """Test that no files are returned when the department directory does not exist."""
        files = report_generator_tool._find_relevant_files(
            department="hr",
            time_dimension="2024",
            report_type="some_report"
        )
        assert len(files) == 0

    def test_read_file_content_success_txt(self, report_generator_tool, workspace_manager):
        """Test reading .txt file content successfully."""
        file_path = workspace_manager.root / "finance" / "balance-sheet-2023.txt"
        content = report_generator_tool._read_file_content(str(file_path))
        assert content == "Assets: 5000, Liabilities: 2000"

    def test_read_file_content_error(self, report_generator_tool):
        """Test handling of file read errors."""
        content = report_generator_tool._read_file_content("/non/existent/file.txt")
        assert "Error reading file" in content

    def test_run_impl_success_csv(self, report_generator_tool, llm_client):
        """Test the full run implementation with a CSV file."""
        tool_input = {
            "report_type": "income_statement",
            "time_dimension": "2024-Q2",
            "department": "finance"
        }
        
        result = report_generator_tool.run_impl(tool_input)

        assert isinstance(result, ToolImplOutput)
        assert result.tool_output == "This is the generated report."
        assert result.tool_result_message == "Report generated successfully."
        llm_client.chat_completion.assert_called_once()
        call_args = llm_client.chat_completion.call_args[1]
        assert "Date,Revenue,Expenses" in call_args['messages'][0]['content']

    def test_run_impl_no_data_found(self, report_generator_tool, llm_client):
        """Test run implementation when no relevant data is found."""
        tool_input = {
            "report_type": "non_existent",
            "time_dimension": "2099",
            "department": "finance"
        }

        result = report_generator_tool.run_impl(tool_input)

        assert isinstance(result, ToolImplOutput)
        assert result.tool_output == "No data found for finance non_existent in 2099."
        assert result.tool_result_message == "Report generation failed."
        llm_client.chat_completion.assert_not_called()

    def test_run_impl_llm_failure(self, report_generator_tool, llm_client):
        """Test run implementation when the LLM call fails."""
        tool_input = {
            "report_type": "income_statement",
            "time_dimension": "2024-Q2",
            "department": "finance"
        }

        llm_client.chat_completion.side_effect = Exception("LLM API Error")

        result = report_generator_tool.run_impl(tool_input)

        assert isinstance(result, ToolImplOutput)
        assert result.tool_output == ""
        assert "Failed to generate report from LLM: LLM API Error" in result.tool_result_message
